import os
import re
from datetime import datetime
from typing import Any, Dict, Iterable, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Xuất dữ liệu OCR ra workbook có cấu trúc và dấu vết kiểm tra."""

    PLACEHOLDER_REGEX = re.compile(r"\{\{([a-zA-Z0-9_\.]+)\}\}")
    NUMBER_TYPES = {"number", "integer", "decimal", "currency", "amount", "percentage"}
    DATE_TYPES = {"date", "datetime"}

    @classmethod
    def export(cls, template_path: str, output_path: str, context: Dict[str, Any]) -> str:
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Excel template not found: {template_path}")
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        keep_vba = template_path.lower().endswith(".xlsm")
        wb = openpyxl.load_workbook(template_path, keep_vba=keep_vba)
        for sheet in wb.worksheets:
            cls._fill_sheet(sheet, context)
        wb.save(output_path)
        wb.close()
        return output_path

    @classmethod
    def export_full_document(cls, document_data: Dict[str, Any], output_path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb = openpyxl.Workbook()
        styles = cls._styles()
        details = document_data.get("field_details") or {}
        metadata = document_data.get("metadata") or {}

        ws_fields = wb.active
        ws_fields.title = "Thông tin bóc tách"
        ws_fields["A1"] = "DỮ LIỆU BÓC TÁCH OCR"
        ws_fields["A1"].font = styles["title"]
        filename = metadata.get("filename") or document_data.get("filename") or ""
        ws_fields["A2"] = f"Tài liệu: {filename} • Tổng số trang: {len(document_data.get('pages', []))}"
        ws_fields["A2"].font = styles["note"]
        ws_fields.append([])
        ws_fields.append(["STT", "Tên trường", "Giá trị bóc tách", "Kiểu dữ liệu", "Độ tin cậy", "Trạng thái"])
        cls._format_header(ws_fields, 4, 6, styles)

        for index, (name, value) in enumerate((document_data.get("fields") or {}).items(), 1):
            detail = details.get(name) or {}
            data_type = str(detail.get("data_type") or "string").lower()
            confidence = cls._confidence(detail.get("confidence"))
            validated = bool(detail.get("validated", confidence >= 0.85))
            error = detail.get("validation_error")
            typed_value = cls._typed_value(value, data_type)
            ws_fields.append([
                index,
                detail.get("label") or name.replace("_", " ").title(),
                typed_value,
                data_type,
                confidence,
                error or ("Đã kiểm tra" if validated else "Cần kiểm tra"),
            ])
            row = ws_fields.max_row
            cls._format_body_row(ws_fields, row, 6, styles)
            ws_fields.cell(row, 5).number_format = "0.0%"
            cls._apply_data_format(ws_fields.cell(row, 3), data_type)
        cls._set_widths(ws_fields, [8, 28, 50, 16, 15, 32])
        ws_fields.freeze_panes = "A5"
        ws_fields.auto_filter.ref = f"A4:F{max(4, ws_fields.max_row)}"

        cls._add_blocks_sheet(wb, document_data.get("pages") or [], styles)
        tables = list(cls._iter_tables(document_data))
        for index, table in enumerate(tables, 1):
            cls._add_table_sheet(wb, table, index, styles)
        cls._add_quality_sheet(wb, document_data, details, tables, styles)

        wb.save(output_path)
        wb.close()
        return output_path

    @classmethod
    def _add_blocks_sheet(cls, wb, pages: List[Dict[str, Any]], styles):
        ws = wb.create_sheet("Đoạn văn bản OCR")
        ws.append(["Trang", "Khối #", "Phân loại", "Nội dung", "Độ tin cậy"])
        cls._format_header(ws, 1, 5, styles)
        for page in pages:
            for idx, block in enumerate(page.get("blocks") or [], 1):
                text = str(block.get("text") or "").strip()
                if not text:
                    continue
                ws.append([page.get("page_number", 1), idx, block.get("tag", "Text"), text,
                           cls._confidence(block.get("confidence"))])
                cls._format_body_row(ws, ws.max_row, 5, styles)
                ws.cell(ws.max_row, 4).alignment = Alignment(vertical="top", wrap_text=True)
                ws.cell(ws.max_row, 5).number_format = "0.0%"
        cls._set_widths(ws, [10, 10, 14, 80, 15])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:E{max(1, ws.max_row)}"

    @classmethod
    def _add_table_sheet(cls, wb, table: Dict[str, Any], index: int, styles):
        title = cls._safe_sheet_title(table.get("name") or f"Bảng {index}", wb.sheetnames)
        ws = wb.create_sheet(title)
        headers = list(table.get("headers") or [])
        rows = list(table.get("rows") or [])
        width = max([len(headers)] + [len(row) for row in rows] + [1])
        if not headers:
            headers = [f"Cột {i}" for i in range(1, width + 1)]
        headers += [f"Cột {i}" for i in range(len(headers) + 1, width + 1)]
        ws.append(headers)
        cls._format_header(ws, 1, width, styles)
        for row in rows:
            values = [cls._infer_cell_value(value) for value in list(row) + [""] * (width - len(row))]
            ws.append(values)
            cls._format_body_row(ws, ws.max_row, width, styles)
        for column in range(1, width + 1):
            longest = max((len(str(ws.cell(r, column).value or "")) for r in range(1, ws.max_row + 1)), default=10)
            ws.column_dimensions[get_column_letter(column)].width = min(max(12, longest + 2), 40)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(width)}{max(1, ws.max_row)}"

    @classmethod
    def _add_quality_sheet(cls, wb, document_data, details, tables, styles):
        ws = wb.create_sheet("Kiểm tra chất lượng")
        ws.append(["Loại", "Vị trí", "Giá trị", "Độ tin cậy", "Kết quả kiểm tra"])
        cls._format_header(ws, 1, 5, styles)
        for name, detail in details.items():
            conf = cls._confidence(detail.get("confidence"))
            error = detail.get("validation_error")
            if conf < 0.85 or error or not detail.get("validated", True):
                ws.append(["Trường", name, detail.get("value", ""), conf, error or "Cần đối chiếu ảnh gốc"])
        for table_idx, table in enumerate(tables, 1):
            table_conf = cls._confidence(table.get("confidence"))
            if table_conf < 0.85:
                ws.append(["Bảng", table.get("name") or f"Bảng {table_idx}", "", table_conf,
                           "Cần đối chiếu cấu trúc bảng"])
            for cell in table.get("cells") or []:
                conf = cls._confidence(cell.get("confidence"))
                if str(cell.get("text") or "").strip() and conf < 0.85:
                    row_index = cell.get("row_index", cell.get("row", 0))
                    col_index = cell.get("col_index", cell.get("col", 0))
                    pos = f"Bảng {table_idx} - R{int(row_index) + 1}C{int(col_index) + 1}"
                    ws.append(["Ô", pos, cell.get("text", ""), conf, "Cần đối chiếu ảnh gốc"])
        for error in (document_data.get("extra") or {}).get("validation_errors", []):
            ws.append(["Quy tắc", "Tài liệu", "", 0, str(error)])
        if ws.max_row == 1:
            ws.append(["Tổng thể", "Tài liệu", "", cls._confidence(document_data.get("confidence")),
                       "Không phát hiện mục dưới ngưỡng; vẫn nên duyệt trước khi phát hành"])
        for row in range(2, ws.max_row + 1):
            cls._format_body_row(ws, row, 5, styles)
            ws.cell(row, 4).number_format = "0.0%"
        cls._set_widths(ws, [14, 32, 55, 15, 42])
        ws.freeze_panes = "A2"

    @staticmethod
    def _iter_tables(document_data: Dict[str, Any]) -> Iterable[Dict[str, Any]]:
        tables = document_data.get("tables") or []
        if tables:
            yield from tables
            return
        for page in document_data.get("pages") or []:
            yield from page.get("tables") or []

    @classmethod
    def _fill_sheet(cls, sheet, context: Dict[str, Any]):
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                original = cell.value
                for key in cls.PLACEHOLDER_REGEX.findall(original):
                    replacement = context.get(key, "")
                    if original.strip() == f"{{{{{key}}}}}":
                        cell.value = cls._infer_cell_value(replacement)
                        break
                    original = original.replace(f"{{{{{key}}}}}", str(replacement or ""))
                    cell.value = original

    @staticmethod
    def _infer_cell_value(value: Any) -> Any:
        if not isinstance(value, str):
            return value
        text = value.strip()
        if not text:
            return ""
        # Không ép mã hồ sơ, số văn bản hoặc chuỗi có số 0 đầu thành số.
        if re.fullmatch(r"[-+]?\d{1,3}(?:[.,]\d{3})+(?:[.,]\d+)?", text):
            normalized = text.replace(".", "").replace(",", ".")
            try:
                return float(normalized)
            except ValueError:
                return value
        if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", text) and not (len(text) > 1 and text.startswith("0")):
            try:
                return float(text.replace(",", "."))
            except ValueError:
                return value
        return value

    @classmethod
    def _typed_value(cls, value: Any, data_type: str) -> Any:
        if value is None:
            return ""
        if data_type in cls.NUMBER_TYPES:
            return cls._infer_cell_value(value)
        if data_type in cls.DATE_TYPES and isinstance(value, str):
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    return datetime.strptime(value.strip(), fmt)
                except ValueError:
                    pass
        return value

    @staticmethod
    def _confidence(value: Any) -> float:
        try:
            result = float(value)
            return max(0.0, min(1.0, result))
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _apply_data_format(cell, data_type: str):
        if data_type in {"currency", "amount", "number", "decimal"}:
            cell.number_format = '#,##0.00'
        elif data_type == "integer":
            cell.number_format = '#,##0'
        elif data_type == "percentage":
            cell.number_format = '0.00%'
        elif data_type in {"date", "datetime"}:
            cell.number_format = 'dd/mm/yyyy'

    @staticmethod
    def _safe_sheet_title(title: str, existing: List[str]) -> str:
        base = re.sub(r"[\\/*?:\[\]]", "_", str(title)).strip()[:31] or "Bảng"
        candidate, suffix = base, 2
        while candidate in existing:
            tail = f" ({suffix})"
            candidate = base[:31 - len(tail)] + tail
            suffix += 1
        return candidate

    @staticmethod
    def _styles():
        side = Side(style="thin", color="CBD5E1")
        border = Border(left=side, right=side, top=side, bottom=side)
        return {
            "header_fill": PatternFill("solid", fgColor="1D4ED8"),
            "header_font": Font(name="Segoe UI", size=10, bold=True, color="FFFFFF"),
            "title": Font(name="Segoe UI", size=14, bold=True, color="0F172A"),
            "note": Font(name="Segoe UI", size=10, italic=True, color="64748B"),
            "normal": Font(name="Segoe UI", size=10),
            "border": border,
        }

    @staticmethod
    def _format_header(ws, row: int, columns: int, styles):
        for col in range(1, columns + 1):
            cell = ws.cell(row, col)
            cell.fill = styles["header_fill"]
            cell.font = styles["header_font"]
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @staticmethod
    def _format_body_row(ws, row: int, columns: int, styles):
        for col in range(1, columns + 1):
            cell = ws.cell(row, col)
            cell.font = styles["normal"]
            cell.border = styles["border"]
            cell.alignment = Alignment(vertical="top")

    @staticmethod
    def _set_widths(ws, widths: List[int]):
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
