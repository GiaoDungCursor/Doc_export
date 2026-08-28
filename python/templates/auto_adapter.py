import re
from pathlib import Path
from typing import Dict, List

from docx import Document


class WordTemplateAutoAdapter:
    """Convert common Vietnamese dotted-line Word forms into mapped templates."""

    PARAGRAPH_RULES = (
        (r"^-\s*Họ tên\b", "- Họ tên: {{full_name}}", ["full_name"]),
        (r"^-\s*Sinh ngày", "- Sinh ngày, tháng, năm: {{birth_date}}    Giới tính: {{gender}}", ["birth_date", "gender"]),
        (r"^-\s*Quê quán", "- Quê quán: {{hometown}}", ["hometown"]),
        (r"^-\s*Trú quán", "- Trú quán: {{residence}}", ["residence"]),
        (r"^-\s*Đơn vị công tác", "- Đơn vị công tác: {{organization}}", ["organization"]),
        (r"^-\s*Chức vụ", "- Chức vụ (Đảng, chính quyền, đoàn thể): {{position}}", ["position"]),
        (r"^-\s*Trình độ chuyên môn", "- Trình độ chuyên môn, nghiệp vụ: {{professional_qualification}}", ["professional_qualification"]),
        (r"^-\s*Học hàm", "- Học hàm, học vị, danh hiệu, giải thưởng: {{academic_qualification}}", ["academic_qualification"]),
        (r"^1\.\s*Quyền hạn, nhiệm vụ", "1. Quyền hạn, nhiệm vụ được giao hoặc đảm nhận: {{assigned_duties}}", ["assigned_duties"]),
        (r"^2\.\s*Thành tích đạt được", "2. Thành tích đạt được của cá nhân: {{achievements}}", ["achievements"]),
    )

    @classmethod
    def adapt(cls, path: str) -> Dict:
        doc = Document(path)
        all_text = "\n".join(p.text for p in doc.paragraphs)
        form_score = len(re.findall(r"\.{5,}|…{3,}", all_text))
        if form_score < 3:
            return {"adapted": False, "adaptations": [], "reason": "not_a_dotted_form"}

        adaptations: List[Dict] = []
        for paragraph in doc.paragraphs:
            current = re.sub(r"\s+", " ", paragraph.text).strip()
            for pattern, replacement, fields in cls.PARAGRAPH_RULES:
                if re.match(pattern, current, re.IGNORECASE):
                    cls._set_text(paragraph, replacement)
                    adaptations.append({"location": current[:80], "fields": fields})
                    break

        # Administrative date row commonly sits in the second row of the first table.
        if doc.tables:
            for row in doc.tables[0].rows:
                for cell in row.cells:
                    if re.search(r"(?i)ngày.*tháng.*năm", cell.text):
                        cls._set_text(cell.paragraphs[0], "{{place}}, ngày {{day}} tháng {{month}} năm {{year}}")
                        adaptations.append({"location": "dòng địa danh, ngày tháng", "fields": ["place", "day", "month", "year"]})

        # Blank data rows under the two award-history tables.
        history_fields = [
            ("commendation_year", "commendation_title", "commendation_decision"),
            ("award_year", "award_form", "award_decision"),
        ]
        for table, fields in zip(doc.tables[1:3], history_fields):
            if len(table.rows) < 2 or len(table.rows[1].cells) < 3:
                continue
            for cell, field in zip(table.rows[1].cells, fields):
                cls._set_text(cell.paragraphs[0], "{{" + field + "}}")
            adaptations.append({"location": "bảng lịch sử khen thưởng", "fields": list(fields)})

        if not adaptations:
            return {"adapted": False, "adaptations": [], "reason": "no_supported_labels"}
        doc.core_properties.comments = "Auto-adapted by Office Studio AI from dotted Word form"
        doc.save(path)
        return {"adapted": True, "adaptations": adaptations, "source_name": Path(path).name}

    @staticmethod
    def _set_text(paragraph, value: str):
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(value)
