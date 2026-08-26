import docx
from typing import Dict, Any, List
import re

class WordTemplateParser:
    """
    Parses Word docx template placeholders such as {{field_name}} in paragraphs & tables
    """

    PLACEHOLDER_REGEX = re.compile(r'\{\{([a-zA-Z0-9_\.]+)\}\}')

    @classmethod
    def find_placeholders(cls, template_path: str) -> List[str]:
        doc = docx.Document(template_path)
        placeholders = set()

        for p in doc.paragraphs:
            if p.text:
                for m in cls.PLACEHOLDER_REGEX.findall(p.text):
                    placeholders.add(m)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for m in cls.PLACEHOLDER_REGEX.findall(p.text):
                            placeholders.add(m)

        return sorted(list(placeholders))
