"""Generate a starter catalog of Vietnamese Office templates with parseable placeholders."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Mm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "app-data" / "templates" / "vietnam"


def paragraph(doc, text="", bold=False, centered=False, size=13):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    return p


def administrative_template(filename, document_name, body_label="", *, parent=False,
                            urgency=False, subject_prefix="", sections=()):
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin, section.right_margin = Mm(30), Mm(20)

    header = doc.add_table(rows=1, cols=2)
    header.autofit = True
    left, right = header.rows[0].cells
    left.text = (("{{parent_authority}}\n" if parent else "")
                 + "{{issuing_authority}}\nSố: {{document_number}}")
    right.text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n{{place}}, ngày {{day}} tháng {{month}} năm {{year}}"
    for cell in (left, right):
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

    if urgency:
        paragraph(doc, "{{urgency}}", True, False, 13)

    paragraph(doc, document_name, True, True, 14)
    paragraph(doc, subject_prefix + "{{title}}", True, True, 13)
    paragraph(doc, "Kính gửi: {{recipient}}")
    paragraph(doc, "{{summary}}")
    for heading, field in sections:
        paragraph(doc, heading, True)
        paragraph(doc, "{{" + field + "}}")
    if not sections:
        paragraph(doc, "{{content}}")
    if body_label:
        paragraph(doc, body_label + ": {{conclusion}}")

    footer = doc.add_table(rows=1, cols=2)
    footer.cell(0, 0).text = "Nơi nhận:\n- {{recipient}};\n- Lưu: {{archive_code}}."
    footer.cell(0, 1).text = "{{signer_title}}\n\n\n{{signer_name}}"
    for cell in footer.rows[0].cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.subject = "vietnamese-administrative:" + filename.removesuffix(".docx")
    doc.save(OUTPUT / filename)


def minutes_template():
    doc = Document()
    paragraph(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", True, True)
    paragraph(doc, "Độc lập - Tự do - Hạnh phúc", True, True)
    paragraph(doc, "BIÊN BẢN", True, True, 15)
    paragraph(doc, "{{title}}", True, True)
    paragraph(doc, "Thời gian: {{date}} - {{time}}")
    paragraph(doc, "Địa điểm: {{place}}")
    paragraph(doc, "Thành phần: {{participants}}")
    paragraph(doc, "Nội dung: {{content}}")
    paragraph(doc, "Kết luận: {{conclusion}}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "THƯ KÝ"
    table.cell(0, 1).text = "CHỦ TRÌ"
    table.cell(1, 0).text = "{{secretary}}"
    table.cell(1, 1).text = "{{chairperson}}"
    doc.core_properties.subject = "vietnamese-administrative:minutes"
    doc.save(OUTPUT / "bien_ban.docx")


def academic_report_template(filename="bao_cao_hoc_thuat.docx"):
    """Neutral academic structure; institution branding remains data, not hard-coded."""
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.top_margin, section.bottom_margin = Mm(25), Mm(25)
    section.left_margin, section.right_margin = Mm(35), Mm(20)
    paragraph(doc, "{{parent_authority}}", True, True, 13)
    paragraph(doc, "{{issuing_authority}}", True, True, 13)
    paragraph(doc, "BÁO CÁO HỌC THUẬT", True, True, 16)
    paragraph(doc, "{{title}}", True, True, 15)
    paragraph(doc, "Tác giả: {{author}}", centered=True)
    paragraph(doc, "Người hướng dẫn: {{supervisor}}", centered=True)
    paragraph(doc, "{{place}}, năm {{year}}", centered=True)
    doc.add_page_break()
    for heading, field in (("TÓM TẮT", "summary"), ("1. GIỚI THIỆU", "introduction"),
                           ("2. PHƯƠNG PHÁP", "methodology"), ("3. KẾT QUẢ", "results"),
                           ("4. KẾT LUẬN VÀ KIẾN NGHỊ", "conclusion"),
                           ("TÀI LIỆU THAM KHẢO", "references")):
        paragraph(doc, heading, True)
        paragraph(doc, "{{" + field + "}}")
    doc.core_properties.subject = "vietnamese-academic-report"
    doc.save(OUTPUT / filename)


def invoice_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Hóa đơn"
    ws.merge_cells("A1:G1")
    ws["A1"] = "HÓA ĐƠN GIÁ TRỊ GIA TĂNG"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    rows = [
        ("A3", "Đơn vị bán: {{supplier}}"), ("A4", "Mã số thuế: {{tax_code}}"),
        ("A5", "Địa chỉ: {{seller_address}}"), ("E3", "Số: {{invoice_number}}"),
        ("E4", "Ngày: {{date}}"), ("A7", "Người mua: {{customer}}"),
        ("A8", "MST người mua: {{buyer_tax_code}}"), ("A9", "Địa chỉ: {{buyer_address}}"),
    ]
    for cell, value in rows:
        ws[cell] = value
    headers = ["STT", "Tên hàng hóa, dịch vụ", "ĐVT", "Số lượng", "Đơn giá", "Thuế suất", "Thành tiền"]
    for col, value in enumerate(headers, 1):
        c = ws.cell(11, col, value)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in range(12, 18):
        for col in range(1, 8):
            ws.cell(row, col, "{{items}}" if row == 12 and col == 2 else "")
    ws["F19"], ws["G19"] = "Cộng tiền hàng", "{{subtotal}}"
    ws["F20"], ws["G20"] = "Tiền thuế GTGT", "{{tax_amount}}"
    ws["F21"], ws["G21"] = "Tổng thanh toán", "{{total_amount}}"
    ws.column_dimensions["B"].width = 32
    wb.properties.subject = "vietnamese-finance:invoice"
    wb.save(OUTPUT / "hoa_don_gtgt.xlsx")


def receipt_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Phiếu thu"
    ws.merge_cells("A1:F1")
    ws["A1"] = "PHIẾU THU"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    values = [
        "Đơn vị: {{issuing_authority}}", "Số phiếu: {{receipt_number}}", "Ngày: {{date}}",
        "Họ tên người nộp tiền: {{payer}}", "Địa chỉ: {{address}}", "Lý do nộp: {{reason}}",
        "Số tiền: {{total_amount}}", "Bằng chữ: {{amount_in_words}}", "Kèm theo: {{attachments}}",
    ]
    for row, value in enumerate(values, 3):
        ws.cell(row, 1, value)
    for col, label in enumerate(["Người lập phiếu", "Người nộp tiền", "Thủ quỹ", "Kế toán trưởng", "Giám đốc"], 1):
        ws.cell(14, col, label).font = Font(bold=True)
    wb.properties.subject = "vietnamese-finance:receipt"
    wb.save(OUTPUT / "phieu_thu.xlsx")


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    administrative_template("cong_van.docx", "")
    administrative_template("cong_van_co_quan_hai_cap.docx", "", parent=True)
    administrative_template("cong_van_khan.docx", "", urgency=True)
    administrative_template("thong_bao.docx", "THÔNG BÁO")
    administrative_template("ke_hoach.docx", "KẾ HOẠCH", sections=(("I. MỤC ĐÍCH, YÊU CẦU", "objectives"), ("II. NỘI DUNG", "content"), ("III. TỔ CHỨC THỰC HIỆN", "implementation")))
    administrative_template("quyet_dinh.docx", "QUYẾT ĐỊNH", "Điều khoản thi hành")
    administrative_template("bao_cao.docx", "BÁO CÁO", "Kiến nghị")
    administrative_template("bao_cao_dinh_ky.docx", "BÁO CÁO", "Kiến nghị", sections=(("I. TÌNH HÌNH, KẾT QUẢ", "content"), ("II. KHÓ KHĂN, VƯỚNG MẮC", "difficulties"), ("III. NHIỆM VỤ, GIẢI PHÁP", "solutions")))
    administrative_template("bao_cao_chuyen_de.docx", "BÁO CÁO CHUYÊN ĐỀ", "Kiến nghị", sections=(("I. BỐI CẢNH", "introduction"), ("II. KẾT QUẢ PHÂN TÍCH", "results"), ("III. KẾT LUẬN", "conclusion")))
    administrative_template("to_trinh.docx", "TỜ TRÌNH", "Kính trình")
    administrative_template("giay_moi.docx", "GIẤY MỜI")
    minutes_template()
    academic_report_template()
    invoice_template()
    receipt_template()


if __name__ == "__main__":
    main()
