import openpyxl
from typing import Dict, Any, List
import re

class ExcelTemplateParser:
    """
    Parses Excel template placeholders such as {{field_name}} or table column tags
    """

    PLACEHOLDER_REGEX = re.compile(r'\{\{([a-zA-Z0-9_\.]+)\}\}')

    @classmethod
    def find_placeholders(cls, template_path: str) -> List[str]:
        wb = openpyxl.load_workbook(template_path, data_only=False)
        placeholders = set()

        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=False):
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        matches = cls.PLACEHOLDER_REGEX.findall(cell.value)
                        for m in matches:
                            placeholders.add(m)

        wb.close()
        return sorted(list(placeholders))
