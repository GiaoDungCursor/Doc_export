import os
import tempfile
import unittest

import cv2
import numpy as np
import openpyxl

from core.document import Block, BoundingBox
from exporters.excel import ExcelExporter
from ocr.table import TableDetector


class TableExcelPipelineTest(unittest.TestCase):
    def test_ruled_table_is_recovered_from_image(self):
        image = np.full((260, 620, 3), 255, dtype=np.uint8)
        xs, ys = [20, 220, 420, 600], [20, 90, 160, 240]
        for x in xs:
            cv2.line(image, (x, ys[0]), (x, ys[-1]), (0, 0, 0), 3)
        for y in ys:
            cv2.line(image, (xs[0], y), (xs[-1], y), (0, 0, 0), 3)
        values = [["Tên", "SL", "Thành tiền"], ["Bút", "2", "20.000"], ["Vở", "3", "45.000"]]
        blocks = []
        for row, row_values in enumerate(values):
            for col, value in enumerate(row_values):
                blocks.append(Block(text=value, confidence=.97, bbox=BoundingBox(
                    x0=xs[col] + 15, y0=ys[row] + 15, x1=xs[col + 1] - 15, y1=ys[row + 1] - 15)))
        tables = TableDetector.extract_tables(image, blocks)
        self.assertEqual(1, len(tables))
        self.assertEqual(values[0], tables[0].headers)
        self.assertEqual(values[1:], tables[0].rows)

    def test_excel_contains_real_confidence_typed_cells_and_table(self):
        document = {
            "metadata": {"filename": "bang.png"},
            "confidence": .91,
            "fields": {"invoice_number": "0012", "total_amount": "65.000"},
            "field_details": {
                "invoice_number": {"label": "Số hóa đơn", "value": "0012", "data_type": "string",
                                   "confidence": .72, "validated": False},
                "total_amount": {"label": "Tổng tiền", "value": "65.000", "data_type": "amount",
                                 "confidence": .98, "validated": True},
            },
            "pages": [],
            "tables": [{"name": "Chi tiết", "headers": ["Tên", "SL", "Thành tiền"],
                        "rows": [["Bút", "2", "20.000"], ["Vở", "3", "45.000"]],
                        "confidence": .94, "cells": []}],
            "extra": {"validation_errors": []},
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "result.xlsx")
            ExcelExporter.export_full_document(document, path)
            wb = openpyxl.load_workbook(path, data_only=False)
            self.assertIn("Chi tiết", wb.sheetnames)
            self.assertIn("Kiểm tra chất lượng", wb.sheetnames)
            fields = wb["Thông tin bóc tách"]
            self.assertEqual("0012", fields["C5"].value)
            self.assertEqual(.72, fields["E5"].value)
            self.assertIsInstance(fields["C6"].value, (int, float))
            detail = wb["Chi tiết"]
            self.assertIsInstance(detail["B2"].value, (int, float))
            self.assertIsInstance(detail["C2"].value, (int, float))
            self.assertGreater(wb["Kiểm tra chất lượng"].max_row, 1)
            wb.close()


if __name__ == "__main__":
    unittest.main()
