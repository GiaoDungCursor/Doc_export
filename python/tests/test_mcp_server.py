import os
import tempfile
import unittest

import docx
import openpyxl

import mcp_server


class McpServerTest(unittest.TestCase):
    def test_initialize_and_tool_listing(self):
        initialized = mcp_server.handle_request({
            "jsonrpc": "2.0", "id": 1, "method": "initialize",
            "params": {"protocolVersion": "2024-11-05"},
        })
        self.assertEqual("office-studio-ai", initialized["result"]["serverInfo"]["name"])
        listed = mcp_server.handle_request({"jsonrpc": "2.0", "id": 2, "method": "tools/list"})
        names = {tool["name"] for tool in listed["result"]["tools"]}
        self.assertIn("ocr_extract_document", names)
        self.assertIn("ocr_map_export", names)
        schema = next(tool for tool in listed["result"]["tools"] if tool["name"] == "ocr_map_export")
        self.assertIn("vision_text", schema["inputSchema"]["properties"])
        self.assertIn("vision_tables", schema["inputSchema"]["properties"])

    def test_vietnamese_official_document_recommends_official_template(self):
        document = {
            "document_type": "generic",
            "raw_text": "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nSố: 12/TANDTC-PC\nV/v thông báo\nKính gửi:",
        }
        selected, _ = mcp_server.recommend_template(document, "word")
        self.assertIsNotNone(selected)
        self.assertIn("cong van", mcp_server._normalize(selected["name"]))

    def test_extracts_docx_text_and_table(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "input.docx")
            document = docx.Document()
            document.add_heading("Báo cáo tháng", level=1)
            document.add_paragraph("Nội dung tiếng Việt")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Tên"
            table.cell(0, 1).text = "Số lượng"
            table.cell(1, 0).text = "Bút"
            table.cell(1, 1).text = "2"
            document.save(path)

            result = mcp_server.tool_ocr_extract_document({"source_path": path})

            self.assertIn("Nội dung tiếng Việt", result["raw_text"])
            self.assertEqual("Tên", result["tables"][0]["headers"][0])
            self.assertEqual("native_docx", result["pages"][0]["extraction_method"])

    def test_extracts_xlsx_sheets_as_tables(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "input.xlsx")
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Dữ liệu"
            sheet.append(["Tên", "Số lượng"])
            sheet.append(["Vở", 3])
            workbook.save(path)
            workbook.close()

            result = mcp_server.tool_ocr_extract_document({"source_path": path})

            self.assertEqual("Dữ liệu", result["tables"][0]["name"])
            self.assertEqual(["Vở", "3"], result["tables"][0]["rows"][0])
            self.assertEqual("native_excel", result["pages"][0]["extraction_method"])


if __name__ == "__main__":
    unittest.main()
